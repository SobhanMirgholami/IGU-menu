from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.shapes import GraphicalProperties
import sys

OUT = sys.argv[1]

BRAND = '0B3D91'
BRAND_SOFT = 'DCE6FF'
POP = 'FFC93C'
INK = '0E1630'
MONEY = '#,##0.00 "₺"'

LABELS = {
    'Türkçe': dict(title='IGU Café · Satış raporu', period='Dönem', summary='Özet', income='Gelir',
                   orders='Sipariş sayısı', average='Ortalama sipariş', cancelled='İptal edilen',
                   hourly='Saatlik satışlar', hour='Saat', count='Sipariş', items='Satılan ürünler',
                   item='Ürün', qty='Adet', chart1='Saatlere göre siparişler', chart2='En çok satan ürünler',
                   rtl=False),
    'English': dict(title='IGU Café · Sales report', period='Period', summary='Summary', income='Income',
                    orders='Orders', average='Average order', cancelled='Cancelled',
                    hourly='Sales by hour', hour='Hour', count='Orders', items='Items sold',
                    item='Item', qty='Qty', chart1='Orders by hour', chart2='Best-selling items',
                    rtl=False),
    'العربية': dict(title='IGU Café · تقرير المبيعات', period='الفترة', summary='الملخص', income='الدخل',
                    orders='عدد الطلبات', average='متوسط الطلب', cancelled='الملغاة',
                    hourly='المبيعات حسب الساعة', hour='الساعة', count='الطلبات', items='المنتجات المباعة',
                    item='المنتج', qty='الكمية', chart1='الطلبات حسب الساعة', chart2='المنتجات الأكثر مبيعًا',
                    rtl=True),
}

HOUR_FIRST, HOUR_LAST = 12, 35      # 24 rows, 00:00 .. 23:00
ITEM_FIRST, ITEM_LAST = 39, 78      # up to 40 items

thin = Side(style='thin', color='C9D2EA')
cell_border = Border(top=thin, bottom=thin, left=thin, right=thin)
base_font = Font(name='Arial', size=10, color=INK)
bold_font = Font(name='Arial', size=10, bold=True, color=INK)
section_font = Font(name='Arial', size=12, bold=True, color=BRAND)
header_fill = PatternFill('solid', fgColor=BRAND_SOFT)

wb = Workbook()
wb.remove(wb.active)

for sheet_name, L in LABELS.items():
    ws = wb.create_sheet(sheet_name)
    ws.sheet_view.rightToLeft = L['rtl']
    ws.sheet_view.showGridLines = False
    ws.column_dimensions['A'].width = 28
    ws.column_dimensions['B'].width = 14
    ws.column_dimensions['C'].width = 16
    ws.column_dimensions['D'].width = 3

    # Title and period
    ws.merge_cells('A1:C1')
    ws['A1'] = L['title']
    ws['A1'].font = Font(name='Arial', size=14, bold=True, color='FFFFFF')
    ws['A1'].fill = PatternFill('solid', fgColor=BRAND)
    ws['A1'].alignment = Alignment(horizontal='left', vertical='center', indent=1)
    ws.row_dimensions[1].height = 28

    ws['A2'] = L['period']
    ws['A2'].font = bold_font
    ws.merge_cells('B2:C2')
    ws['B2'].font = base_font          # filled by the app
    ws['B2'].border = cell_border

    # Summary: totals are formulas over the hourly table
    ws['A4'] = L['summary']
    ws['A4'].font = section_font
    summary = [
        (5, L['income'], f'=SUM(C{HOUR_FIRST}:C{HOUR_LAST})', MONEY),
        (6, L['orders'], f'=SUM(B{HOUR_FIRST}:B{HOUR_LAST})', '0'),
        (7, L['average'], '=IFERROR(B5/B6,0)', MONEY),
        (8, L['cancelled'], 0, '0'),    # filled by the app
    ]
    for row, label, value, fmt in summary:
        ws[f'A{row}'] = label
        ws[f'A{row}'].font = base_font
        ws[f'A{row}'].border = cell_border
        ws[f'B{row}'] = value
        ws[f'B{row}'].number_format = fmt
        ws[f'B{row}'].font = bold_font
        ws[f'B{row}'].border = cell_border

    # Hourly table
    ws['A10'] = L['hourly']
    ws['A10'].font = section_font
    for col, text in zip('ABC', [L['hour'], L['count'], L['income']]):
        c = ws[f'{col}11']
        c.value = text
        c.font = bold_font
        c.fill = header_fill
        c.border = cell_border
    for i, row in enumerate(range(HOUR_FIRST, HOUR_LAST + 1)):
        ws[f'A{row}'] = f'{i:02d}:00'
        ws[f'A{row}'].font = base_font
        ws[f'B{row}'] = 0
        ws[f'B{row}'].number_format = '0'
        ws[f'C{row}'] = 0
        ws[f'C{row}'].number_format = MONEY
        for col in 'ABC':
            ws[f'{col}{row}'].border = cell_border
            if col != 'A':
                ws[f'{col}{row}'].font = base_font

    # Items table
    ws['A37'] = L['items']
    ws['A37'].font = section_font
    for col, text in zip('ABC', [L['item'], L['qty'], L['income']]):
        c = ws[f'{col}38']
        c.value = text
        c.font = bold_font
        c.fill = header_fill
        c.border = cell_border
    for row in range(ITEM_FIRST, ITEM_LAST + 1):
        ws[f'A{row}'].font = base_font
        ws[f'B{row}'].number_format = '0'
        ws[f'B{row}'].font = base_font
        ws[f'C{row}'].number_format = MONEY
        ws[f'C{row}'].font = base_font
        for col in 'ABC':
            ws[f'{col}{row}'].border = cell_border

    # Chart 1: orders by hour
    c1 = BarChart()
    c1.type = 'col'
    c1.title = L['chart1']
    c1.legend = None
    c1.y_axis.title = L['count']
    c1.x_axis.title = L['hour']
    c1.y_axis.majorGridlines = None
    c1.x_axis.delete = False
    c1.y_axis.delete = False
    c1.add_data(Reference(ws, min_col=2, min_row=11, max_row=HOUR_LAST), titles_from_data=True)
    c1.set_categories(Reference(ws, min_col=1, min_row=HOUR_FIRST, max_row=HOUR_LAST))
    c1.series[0].graphicalProperties = GraphicalProperties(solidFill=BRAND)
    c1.series[0].graphicalProperties.line.solidFill = INK
    c1.gapWidth = 60
    c1.width, c1.height = 18, 8.5
    ws.add_chart(c1, 'E2')

    # Chart 2: best-selling items
    c2 = BarChart()
    c2.type = 'bar'
    c2.title = L['chart2']
    c2.legend = None
    c2.x_axis.delete = False
    c2.y_axis.delete = False
    c2.y_axis.majorGridlines = None
    c2.x_axis.scaling.orientation = 'maxMin'   # first item on top
    c2.add_data(Reference(ws, min_col=2, min_row=38, max_row=ITEM_LAST), titles_from_data=True)
    c2.set_categories(Reference(ws, min_col=1, min_row=ITEM_FIRST, max_row=ITEM_LAST))
    c2.series[0].graphicalProperties = GraphicalProperties(solidFill=POP)
    c2.series[0].graphicalProperties.line.solidFill = INK
    c2.gapWidth = 50
    c2.width, c2.height = 18, 8.5
    ws.add_chart(c2, 'E20')

    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True

wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print('saved', OUT)
